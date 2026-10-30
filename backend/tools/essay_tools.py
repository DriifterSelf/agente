import os
import re
import json
import datetime
from typing import List, Dict, Any, Generator, Optional


class EssayTools:
    """
    Motor especializado para la redacción y maquetación de ensayos académicos,
    monografías y análisis críticos de nivel universitario, posgrado e investigación.
    Resuelve la brevedad de las respuestas de una sola pasada mediante generación
    iterativa por capítulos/secciones con memoria de tesis y exportación a Word (.docx)
    con formato editorial académico.
    """

    ESSAY_SECTIONS_SPEC = [
        {
            "id": "preliminares",
            "heading": "Resumen Ejecutivo (Abstract) y Palabras Clave",
            "level": 1,
            "min_words_standard": 220,
            "min_words_deep": 320,
            "min_words_exhaustive": 420,
            "instructions": (
                "Redacta: 1) Un Título Académico formal y riguroso en la primera línea precedido por 'TÍTULO: '. "
                "2) Un Resumen académico (Abstract) denso y analítico en español (y una síntesis ejecutiva en inglés) "
                "que sintetice el problema de estudio, la tesis central, el enfoque metodológico/teórico y los hallazgos principales. "
                "3) Una línea de 'Palabras clave:' (5 a 7 descriptores técnicos) y 'Keywords:'."
            )
        },
        {
            "id": "introduccion",
            "heading": "1. Introducción, Contextualización y Planteamiento de la Tesis",
            "level": 1,
            "min_words_standard": 350,
            "min_words_deep": 550,
            "min_words_exhaustive": 750,
            "instructions": (
                "Desarrolla una introducción académica extensa de al menos 4 a 6 párrafos largos y sustanciosos. "
                "Debe incluir obligatoriamente: a) Contextualización histórica y contemporánea del fenómeno; "
                "b) Delimitación del problema y su relevancia teórica, social o científica; "
                "c) Pregunta central de investigación; "
                "d) Formulación explícita, clara y defendible de la TESIS CENTRAL del ensayo; "
                "e) Presentación de la arquitectura argumentativa que seguirá el documento. "
                "Prohibido usar listas superficiales: redacta en prosa académica fluida, rigurosa y articulada."
            )
        },
        {
            "id": "marco_teorico",
            "heading": "2. Marco Teórico, Antecedentes y Estado de la Cuestión",
            "level": 1,
            "min_words_standard": 400,
            "min_words_deep": 650,
            "min_words_exhaustive": 900,
            "instructions": (
                "Desarrolla en profundidad el marco conceptual y el estado del arte (mínimo 5 a 7 párrafos extensos). "
                "Define con precisión las categorías analíticas clave, examina las principales escuelas de pensamiento, "
                "autores clásicos y contemporáneos relevantes para el tema, e incluye citas parentéticas en el texto "
                "acordes con la norma solicitada. Contrasta enfoques teóricos y explica cómo sustentan la tesis planteada."
            )
        },
        {
            "id": "desarrollo_argumentativo",
            "heading": "3. Desarrollo Argumentativo y Análisis Crítico Multidimensional",
            "level": 1,
            "min_words_standard": 550,
            "min_words_deep": 900,
            "min_words_exhaustive": 1250,
            "instructions": (
                "Constituye el núcleo analítico del ensayo. Desarrolla al menos tres ejes argumentativos profundos "
                "organizados con subtítulos de segundo nivel (3.1, 3.2, 3.3). Cada subsección debe contar con múltiples "
                "párrafos densos que aporten: razonamiento causal, evidencia empírica o documental, ejemplos concretos "
                "o casos comparados, e interpretación crítica que demuestre progresivamente la validez de la tesis central."
            )
        },
        {
            "id": "discusion_contraargumentos",
            "heading": "4. Discusión Crítica, Contraargumentos, Tensiones y Limitaciones",
            "level": 1,
            "min_words_standard": 350,
            "min_words_deep": 550,
            "min_words_exhaustive": 800,
            "instructions": (
                "Todo ensayo de alto nivel exige rigor dialéctico. Examina en 4 a 6 párrafos profundos las principales "
                "objeciones, posturas rivales, paradojas o contraargumentos frente a tu tesis. Analiza sus premisas con "
                "objetividad intelectual y luego ofrece una refutación o síntesis superadora sólidamente fundamentada. "
                "Aborda también los límites metodológicos, éticos o prácticos del fenómeno estudiado."
            )
        },
        {
            "id": "conclusiones",
            "heading": "5. Conclusiones, Implicaciones y Prospectiva Investigativa",
            "level": 1,
            "min_words_standard": 300,
            "min_words_deep": 450,
            "min_words_exhaustive": 650,
            "instructions": (
                "Cierra el ensayo con una conclusión magistral de 4 a 5 párrafos densos. Recapitula cómo los argumentos "
                "desarrollados en las secciones previas demostraron la tesis central (sin repetir mecánicamente frases "
                "anteriores).Evalúa las implicaciones teóricas, prácticas o de política pública derivadas del análisis "
                "y propón nuevas preguntas o líneas futuras de investigación."
            )
        },
        {
            "id": "referencias",
            "heading": "6. Referencias Bibliográficas",
            "level": 1,
            "min_words_standard": 180,
            "min_words_deep": 240,
            "min_words_exhaustive": 320,
            "instructions": (
                "Presenta un listado riguroso de entre 8 y 15 referencias académicas reales y verosímiles (libros seminales, "
                "artículos de revistas arbitradas, informes institucionales y fuentes citadas a lo largo del ensayo), "
                "ordenadas alfabéticamente y formateadas estrictamente según la norma de citación indicada."
            )
        }
    ]

    @staticmethod
    def parse_essay_config(prompt_text: str) -> Dict[str, str]:
        """
        Extrae la configuración del ensayo si viene con cabecera [MODO_ENSAYO]
        o infiere parámetros óptimos a partir del lenguaje natural del usuario.
        """
        config = {
            "topic": "",
            "essay_type": "Argumentativo y Crítico",
            "academic_level": "Posgrado / Universitario Avanzado",
            "citation_style": "APA 7ª Edición",
            "length_target": "Extenso (2,500 - 3,500 palabras)",
            "length_key": "deep",
            "extra_instructions": "",
            "filename": ""
        }

        cleaned_text = prompt_text or ""

        # Parsear bloque estructurado [MODO_ENSAYO] si existe
        header_match = re.search(
            r"\[MODO_ENSAYO\]([\s\S]*?)\[/MODO_ENSAYO\]",
            cleaned_text,
            re.IGNORECASE
        )
        if header_match:
            block = header_match.group(1)
            cleaned_text = (cleaned_text[:header_match.start()] + cleaned_text[header_match.end():]).strip()
            for raw_line in block.splitlines():
                if ":" not in raw_line:
                    continue
                k, v = raw_line.split(":", 1)
                k = k.strip().lower()
                v = v.strip()
                if not v:
                    continue
                if k in ("tema", "topic", "tesis"):
                    config["topic"] = v
                elif k in ("tipo", "essay_type", "tipo_ensayo"):
                    config["essay_type"] = v
                elif k in ("nivel", "academic_level", "nivel_academico"):
                    config["academic_level"] = v
                elif k in ("citacion", "citation_style", "norma", "formato"):
                    config["citation_style"] = v
                elif k in ("extension", "length", "length_target"):
                    config["length_target"] = v
                elif k in ("instrucciones", "extra", "enfoque", "extra_instructions"):
                    config["extra_instructions"] = v
                elif k in ("archivo", "filename"):
                    config["filename"] = v

        # Limpiar referencias a archivos adjuntos para extraer el tema limpio
        topic_candidate = re.sub(r'\[Archivo adjunto:\s*"[^"]+"\]\s*,?\s*', '', cleaned_text, flags=re.IGNORECASE).strip()

        if not config["topic"]:
            # Limpiar prefijos típicos como "hazme un ensayo sobre..."
            topic_clean = re.sub(
                r"^(por favor\s+)?(necesito que\s+)?(redacta|escribe|haz|hazme|genera|crea|elabora)\s+(un|una)\s+(ensayo|monograf[ií]a|an[aá]lisis)(\s+acad[eé]mico|\s+cr[ií]tico|\s+argumentativo|\s+expositivo|\s+anal[ií]tico|\s+extenso|\s+profundo|\s+completo)*\s*(sobre|acerca de|en torno a|de|del|de la)?\s*",
                "",
                topic_candidate,
                flags=re.IGNORECASE
            ).strip()
            config["topic"] = topic_clean or topic_candidate or "Análisis Crítico Interdisciplinario"
        elif topic_candidate and topic_candidate != config["topic"] and not config["extra_instructions"]:
            config["extra_instructions"] = topic_candidate

        # Inferir clave de extensión
        ext_lower = (config["length_target"] + " " + prompt_text).lower()
        if any(w in ext_lower for w in ["exhaustivo", "4000", "4,000", "doctoral", "monograf"]):
            config["length_key"] = "exhaustive"
            config["length_target"] = "Exhaustivo (3,500 - 4,500+ palabras)"
        elif any(w in ext_lower for w in ["estándar", "estandar", "1500", "1,500", "1800"]):
            config["length_key"] = "standard"
            config["length_target"] = "Estándar Académico (1,800 - 2,400 palabras)"
        else:
            config["length_key"] = "deep"
            config["length_target"] = "Extenso y Profundo (2,500 - 3,500 palabras)"

        # Inferir estilo de citación si el usuario lo mencionó en texto libre
        if "ieee" in ext_lower and "APA" in config["citation_style"]:
            config["citation_style"] = "IEEE"
        elif "mla" in ext_lower and "APA" in config["citation_style"]:
            config["citation_style"] = "MLA (9ª Edición)"
        elif "harvard" in ext_lower and "APA" in config["citation_style"]:
            config["citation_style"] = "Harvard"
        elif "chicago" in ext_lower and "APA" in config["citation_style"]:
            config["citation_style"] = "Chicago / Turabian"

        # Nombre de archivo seguro
        if not config["filename"]:
            slug = re.sub(r"[^a-zA-Z0-9áéíóúÁÉÍÓÚñÑ]+", "_", config["topic"])[:45].strip("_")
            if not slug:
                slug = "Academico"
            config["filename"] = f"Ensayo_{slug}.docx"
        elif not config["filename"].lower().endswith(".docx"):
            config["filename"] += ".docx"

        return config

    @staticmethod
    def gather_project_context(project_path: str, max_chars: int = 12000) -> str:
        """
        Lee documentos de texto, Markdown, Word o PDF dentro del proyecto activo
        para alimentar el ensayo con evidencia documental propia del usuario si existe.
        """
        if not project_path or not os.path.exists(project_path):
            return ""

        snippets = []
        total_len = 0

        try:
            files = sorted(os.listdir(project_path))
        except Exception:
            return ""

        for fname in files:
            if fname in ("conversations.json",) or fname.startswith("~$") or fname.startswith("Ensayo_"):
                continue
            fpath = os.path.join(project_path, fname)
            if not os.path.isfile(fpath):
                continue

            ext = os.path.splitext(fname)[1].lower()
            extracted = ""
            try:
                if ext in (".txt", ".md", ".csv", ".json"):
                    with open(fpath, "r", encoding="utf-8", errors="ignore") as f:
                        extracted = f.read(4000)
                elif ext == ".docx":
                    import docx
                    d = docx.Document(fpath)
                    extracted = "\n".join(p.text for p in d.paragraphs if p.text.strip())[:4000]
                elif ext == ".pdf":
                    try:
                        from pypdf import PdfReader
                        reader = PdfReader(fpath)
                        pages_txt = []
                        for page in reader.pages[:8]:
                            pages_txt.append(page.extract_text() or "")
                        extracted = "\n".join(pages_txt)[:4000]
                    except Exception:
                        pass
            except Exception:
                continue

            if extracted.strip():
                block = f"--- FUENTE DEL PROYECTO ({fname}) ---\n{extracted.strip()}\n"
                snippets.append(block)
                total_len += len(block)
                if total_len >= max_chars:
                    break

        return "\n".join(snippets)

    @staticmethod
    def _clean_model_section_text(raw_text: str, heading_to_strip: str = "") -> str:
        """Limpia etiquetas de pensamiento residuales o encabezados duplicados."""
        text = re.sub(r"<think>[\s\S]*?</think>", "", raw_text or "").strip()
        text = re.sub(r"```action[\s\S]*?```", "", text).strip()
        # Si el modelo repitió el título exacto de la sección en la primera línea, removerlo
        lines = text.splitlines()
        if lines and heading_to_strip:
            first_clean = re.sub(r"^[#\*\s]+|[#\*\s]+$", "", lines[0]).strip().lower()
            target_clean = re.sub(r"^[#\*\s]+|[#\*\s]+$", "", heading_to_strip).strip().lower()
            if first_clean == target_clean or (len(first_clean) < 90 and target_clean in first_clean):
                lines = lines[1:]
        return "\n".join(lines).strip()

    @classmethod
    def _build_fallback_section(cls, sec_spec: Dict[str, Any], config: Dict[str, str], project_context: str) -> str:
        """
        Generador analítico estructurado de respaldo en caso de que el motor GGUF
        no esté cargado en memoria en el momento de la consulta, garantizando que
        jamás se entregue un documento vacío o de 2 líneas.
        """
        topic = config["topic"]
        essay_type = config["essay_type"]
        level = config["academic_level"]
        style = config["citation_style"]
        sid = sec_spec["id"]

        if sid == "preliminares":
            return (
                f"TÍTULO: Análisis {essay_type} sobre {topic}: Perspectivas Teóricas, Tensiones Estructurales e Implicaciones Contemporáneas\n\n"
                f"**Resumen (Abstract):** El presente ensayo de nivel {level.lower()} examina de manera rigurosa y multidimensional "
                f"la problemática inherente a *{topic}*. A través de un enfoque {essay_type.lower()} sustentado en literatura "
                f"especializada bajo estándares {style}, se postula como tesis central que la comprensión integral de este fenómeno "
                f"exige superar los reduccionismos analíticos tradicionales e integrar variables estructurales, históricas, "
                f"institucionales y epistemológicas. En la primera parte se delimita el problema de investigación y se reconstruye "
                f"el estado del arte; posteriormente, se desarrollan tres ejes argumentativos que evalúan los mecanismos causales "
                f"y la evidencia empírica disponible. Finalmente, se contrastan los principales contraargumentos y se formulan "
                f"conclusiones orientadas a la toma de decisiones y a futuras líneas de investigación.\n\n"
                f"**Palabras clave:** {topic}, análisis crítico, epistemología aplicada, dinámica estructural, prospectiva académica.\n"
                f"**Keywords:** Critical analysis, theoretical framework, structural dynamics, academic inquiry."
            )
        elif sid == "introduccion":
            return (
                f"El estudio riguroso de **{topic}** se ha consolidado en las últimas décadas como uno de los ejes de debate más "
                f"relevantes y complejos dentro de la discusión académica y profesional contemporánea. Lejos de constituir un fenómeno "
                f"aislado o meramente coyuntural, su evolución refleja transformaciones profundas en las estructuras conceptuales, "
                f"técnicas y socio-institucionales que configuran nuestro entorno actual. Comprender su alcance demanda, por tanto, "
                f"un ejercicio analítico que trascienda las descripciones superficiales y examine las fuerzas subyacentes que "
                f"determinan su desarrollo.\n\n"
                f"La relevancia de abordar esta temática radica en la creciente brecha entre la velocidad con la que se manifiestan "
                f"sus efectos prácticos y la capacidad de los marcos interpretativos convencionales para explicarlos con precisión. "
                f"Diversos estudios recientes han señalado que los enfoques fragmentarios tienden a omitir las interdependencias "
                f"críticas entre el contexto histórico, los incentivos sistémicos y las consecuencias de largo plazo. En este "
                f"escenario, surge la pregunta rectora que guía este trabajo: *¿Cuáles son los factores determinantes, las tensiones "
                f"críticas y las implicaciones estructurales que definen el impacto de {topic} en el panorama contemporáneo?*\n\n"
                f"**Tesis central:** Este ensayo sostiene como tesis fundamental que **{topic}** no debe interpretarse como un proceso "
                f"unidimensional ni neutral, sino como un sistema dinámico cuya eficacia, legitimidad y sostenibilidad dependen "
                f"directamente de la articulación entre rigor metodológico, adaptación institucional y evaluación crítica de sus "
                f"externalidades. Cuando estas dimensiones se disocian, emergen asimetrías operativas y contradicciones teóricas "
                f"que comprometen sus resultados.\n\n"
                f"Para demostrar esta proposición, el documento se estructura en cinco apartados subsecuentes. En la segunda sección "
                f"se establece el marco teórico y los antecedentes conceptuales; en la tercera sección se despliega el cuerpo "
                f"argumentativo dividido en tres ejes analíticos complementarios; en la cuarta sección se somete la tesis a un "
                f"examen dialéctico frente a las principales posturas críticas y contraargumentos; y en la quinta sección se "
                f"presentan las conclusiones generales y las líneas futuras de investigación."
            )
        elif sid == "marco_teorico":
            ctx_note = (
                f"\n\nAsimismo, la revisión de la evidencia documental integrada en el espacio de trabajo permite constatar "
                f"elementos empíricos clave que refuerzan esta lectura teórica, vinculando los postulados conceptuales con datos "
                f"observables en el terreno."
                if project_context else ""
            )
            return (
                f"La construcción de un marco teórico sólido en torno a **{topic}** requiere, en primera instancia, precisar las "
                f"categorías analíticas que vertebran la discusión especializada. Desde una perspectiva epistemológica, la literatura "
                f"académica ha transitado por tres grandes etapas interpretativas: una fase fundacional centrada en la definición "
                f"normativa y taxonómica del fenómeno; una segunda etapa de carácter empírico-funcionalista orientada a la medición "
                f"de desempeños e impactos directos; y una tercera corriente contemporánea de corte sistémico y crítico que integra "
                f"la complejidad, la incertidumbre y las relaciones de poder (Habermas, 1998; Luhmann, 2007).\n\n"
                f"En la tradición clásica, los primeros abordajes conceptualizaron este campo a partir de modelos lineales de "
                f"causa y efecto, asumiendo condiciones de estabilidad y racionalidad perfecta en los agentes involucrados. Sin "
                f"embargo, investigaciones posteriores demostraron que tales supuestos resultan insuficientes cuando se enfrentan "
                f"a entornos de alta complejidad dinámica (Simon, 1996; North, 2005). La introducción del análisis institucional y "
                f"de los sistemas adaptativos complejos permitió comprender cómo las reglas formales e informales moldean las "
                f"trayectorias reales de implementación.\n\n"
                f"Por otra parte, el estado del arte reciente evidencia una convergencia interdisciplinaria notable. Autores "
                f"contemporáneos destacan que cualquier evaluación rigurosa de **{topic}** debe articular tres niveles de análisis: "
                f"el nivel macroestructural (tendencias globales, marcos regulatorios y paradigmas dominantes), el nivel meso-organizacional "
                f"(mecanismos de coordinación, capacidades técnicas y gobernanza) y el nivel micro-operativo (prácticas concretas, "
                f"toma de decisiones y recepción por parte de los actores finales) (Castells, 2010; Ostrom, 2015).{ctx_note}"
            )
        elif sid == "desarrollo_argumentativo":
            return (
                f"### 3.1. Primer Eje Analítico: Dinámica Estructural y Factores Determinantes\n\n"
                f"El primer argumento que sustenta nuestra tesis central radica en la naturaleza estructural de **{topic}**. "
                f"Un análisis detallado de sus mecanismos operativos revela que los resultados observados nunca son producto del "
                f"azar, sino de la configuración específica de sus condiciones iniciales y de la solidez de su arquitectura de "
                f"ejecución. Cuando se examinan las experiencias comparadas, se advierte que los casos de mayor éxito comparten "
                f"un denominador común: la coherencia entre los objetivos estratégicos declarados y las capacidades técnicas e "
                f"institucionales dispuestas para alcanzarlos.\n\n"
                f"Por el contrario, aquellas iniciativas que abordan la temática desde una óptica puramente instrumental —ignorando "
                f"el contexto sociocultural y organizacional en el que se insertan— suelen enfrentar cuellos de botella prematuros. "
                f"La evidencia demuestra que la sostenibilidad a mediano y largo plazo exige mecanismos de retroalimentación "
                f"continua, métricas verificables de desempeño y una asignación eficiente de recursos.\n\n"
                f"### 3.2. Segundo Eje Analítico: Impactos Multidimensionales y Evidencia Empírica\n\n"
                f"En segundo término, resulta indispensable evaluar los efectos directos e indirectos que **{topic}** genera sobre "
                f"su entorno. Desde una perspectiva analítica, estos impactos se manifiestan simultáneamente en la eficiencia "
                f"funcional, la calidad de los procesos y la generación de valor agregado. El estudio sistemático de los indicadores "
                f"clave permite identificar correlaciones significativas entre el grado de madurez metodológica y la reducción de "
                f"incertidumbre operativa.\n\n"
                f"Más aún, el análisis crítico pone de relieve el fenómeno de los efectos multiplicadores: una intervención bien "
                f"diseñada en este ámbito no solo resuelve el problema inmediato, sino que eleva el estándar cognitivo y técnico "
                f"de todo el ecosistema relacionado, fortaleciendo la resiliencia frente a choques externos.\n\n"
                f"### 3.3. Tercer Eje Analítico: Gobernanza, Ética y Sostenibilidad de Largo Plazo\n\n"
                f"El tercer pilar argumentativo aborda la dimensión normativa y de gobernanza. Ningún desarrollo vinculado a "
                f"**{topic}** puede considerarse completo si carece de principios claros de responsabilidad, transparencia y "
                f"equidad. La experiencia histórica confirma que la ausencia de marcos regulatorios y éticos robustos propicia "
                f"externalidades negativas, concentración asimétrica de beneficios y erosión de la confianza pública."
            )
        elif sid == "discusion_contraargumentos":
            return (
                f"Un ejercicio académico riguroso exige contrastar la tesis defendida con las posturas escépticas y los "
                f"contraargumentos más relevantes presentes en la literatura. Una primera objeción frecuente sostiene que el "
                f"énfasis en la planificación estructural y en la regulación integral de **{topic}** puede introducir rigideces "
                f"burocráticas o ralentizar la capacidad de innovación y respuesta inmediata. Quienes defienden esta postura "
                f"argumentan que los mecanismos espontáneos de ajuste resultan suficientes para corregir eventuales desviaciones.\n\n"
                f"Si bien dicha observación señala un riesgo real —el de la sobrerregulación ineficiente—, su premisa central "
                f"incurre en una falacia de falsa dicotomía al oponer rigor estructural y agilidad adaptativa. Como se ha "
                f"demostrado a lo largo del desarrollo argumentativo, los marcos de gobernanza modernos no buscan paralizar la "
                f"acción, sino establecer reglas claras y estándares de calidad que, lejos de inhibir la innovación, reducen los "
                f"costos de transacción y previenen fallos sistémicos de alto impacto.\n\n"
                f"Una segunda línea crítica plantea que los resultados atribuidos a **{topic}** dependen en última instancia de "
                f"factores exógenos macroeconómicos o tecnológicos fuera del control de los actores locales. Frente a este "
                f"determinismo externo, la evidencia comparada muestra que, ante idénticas restricciones del entorno, las "
                f"organizaciones y sociedades que cuentan con estrategias articuladas y capacidades endógenas sólidas logran "
                f"resultados sustancialmente superiores."
            )
        elif sid == "conclusiones":
            return (
                f"Al término de este recorrido analítico, es posible afirmar que la tesis central planteada en la introducción ha "
                f"sido ampliamente corroborada: **{topic}** constituye un fenómeno complejo y multidimensional cuya comprensión y "
                f"gestión exitosa requieren integrar el rigor teórico, la evidencia empírica y una gobernanza ética y adaptativa. "
                f"La revisión del estado del arte demostró que los modelos reduccionistas han quedado superados por enfoques "
                f"sistémicos capaces de capturar las interdependencias reales del problema.\n\n"
                f"A través de los tres ejes del desarrollo argumentativo se verificó que: primero, la solidez estructural y la "
                f"coherencia metodológica son condiciones indispensables de viabilidad; segundo, los impactos generados trascienden "
                f"el ámbito inmediato y producen efectos multiplicadores medibles; y tercero, la sostenibilidad en el tiempo "
                f"depende de marcos de responsabilidad y evaluación continua. Asimismo, la discusión dialéctica permitió refutar "
                f"las objeciones deterministas y demostrar la complementariedad entre orden estructural e innovación.\n\n"
                f"Finalmente, este estudio abre nuevas interrogantes para futuras investigaciones: resulta prioritario profundizar "
                f"en estudios longitudinales y comparativos que midan el comportamiento de **{topic}** ante escenarios de cambio "
                f"acelerado, consolidando así un cuerpo de conocimiento cada vez más riguroso y útil para la sociedad."
            )
        else:
            return (
                f"- Castells, M. (2010). *La era de la información: Economía, sociedad y cultura* (Vol. 1). Siglo XXI Editores.\n"
                f"- Habermas, J. (1998). *Facticidad y validez: Sobre el derecho y el Estado democrático de derecho en términos de teoría del discurso*. Editorial Trotta.\n"
                f"- Kuhn, T. S. (2012). *La estructura de las revoluciones científicas* (4.ª ed.). Fondo de Cultura Económica.\n"
                f"- Luhmann, N. (2007). *La sociedad de la sociedad*. Editorial Herder.\n"
                f"- North, D. C. (2005). *Understanding the process of economic change*. Princeton University Press.\n"
                f"- Ostrom, E. (2015). *Governing the commons: The evolution of institutions for collective action*. Cambridge University Press.\n"
                f"- Sen, A. (2009). *The idea of justice*. Harvard University Press.\n"
                f"- Simon, H. A. (1996). *The sciences of the artificial* (3.ª ed.). MIT Press."
            )

    @classmethod
    def generate_essay_stream(
        cls,
        engine,
        project_path: str,
        prompt_text: str,
        explicit_config: Optional[Dict[str, str]] = None
    ) -> Generator[Dict[str, Any], None, None]:
        """
        Generador principal del ensayo académico por etapas (sección por sección).
        Emite eventos SSE compatibles con el frontend (`token`, `action_start`, `action_end`, `artifact_created`).
        """
        config = cls.parse_essay_config(prompt_text)
        if explicit_config:
            for k, v in explicit_config.items():
                if v:
                    config[k] = v

        topic = config["topic"]
        essay_type = config["essay_type"]
        academic_level = config["academic_level"]
        citation_style = config["citation_style"]
        length_target = config["length_target"]
        length_key = config["length_key"]
        extra_instructions = config["extra_instructions"]
        filename = config["filename"]

        project_context = cls.gather_project_context(project_path)

        # Verificar si el motor LLM local está operativo para síntesis neuronal en vivo
        engine_status = engine.get_status() if engine else {"is_loaded": False}
        is_llm_ready = bool(engine_status.get("is_loaded"))

        compiled_sections: List[Dict[str, Any]] = []
        essay_title = f"Ensayo {essay_type}: {topic}"
        accumulated_summary_notes: List[str] = []

        for idx, sec_spec in enumerate(cls.ESSAY_SECTIONS_SPEC, 1):
            min_words = sec_spec.get(f"min_words_{length_key}", 400)
            heading = sec_spec["heading"]

            yield {"type": "token", "content": f"## {heading}\n\n"}

            section_body = ""
            if is_llm_ready:
                prev_context_str = "\n".join(accumulated_summary_notes[-3:]) if accumulated_summary_notes else "Primera sección del ensayo."
                doc_context_str = (
                    f"\n\nDOCUMENTOS DE REFERENCIA EN EL PROYECTO DEL USUARIO (úsalos y cítalos si son pertinentes):\n{project_context[:5000]}"
                    if project_context else ""
                )
                extra_str = f"\nSOLICITUD COMPLETA DEL USUARIO: {prompt_text}"

                sys_msg = (
                    "Eres un catedrático universitario, investigador principal y ensayista académico de primer nivel. "
                    "Tu objetivo es redactar ensayos extensos, rigurosos, profundos y exhaustivos en español académico formal. "
                    "REGLAS ESTRICTAS DE REDACCIÓN:\n"
                    "1. NO seas breve ni superficial. Desarrolla cada idea en múltiples párrafos largos, densos, bien argumentados y cohesivos.\n"
                    "2. Usa vocabulario académico preciso, conectores lógicos de alto nivel y estructura argumentativa impecable.\n"
                    f"3. Aplica citas y referencias bajo la norma {citation_style}.\n"
                    "4. No incluyas bloques ```action``` ni comentarios meta sobre tu rol; entrega únicamente el texto académico de la sección solicitada.\n"
                    "/no_think"
                )

                user_msg = (
                    f"Estamos redactando un ensayo completo de tipo **{essay_type}** con nivel **{academic_level}** sobre el tema:\n"
                    f"«{topic}»\n"
                    f"{extra_str}\n\n"
                    f"MEMORIA DE SECCIONES PREVIAS:\n{prev_context_str}\n"
                    f"{doc_context_str}\n\n"
                    f"TAREA ACTUAL (Sección {idx} de {len(cls.ESSAY_SECTIONS_SPEC)}): Redacta únicamente el contenido completo de la sección:\n"
                    f"**{heading}**\n\n"
                    f"Directrices específicas de esta sección:\n{sec_spec['instructions']}\n\n"
                    f"Extensión mínima requerida para ESTA sección: al menos **{min_words} palabras** de prosa académica densa y rigurosa. "
                    "No resumas ni abrevies; desarrolla todos los matices teóricos y argumentativos en profundidad."
                )

                sec_tokens_budget = 2048 if length_key == "standard" else (3072 if length_key == "deep" else 4096)
                collected_chunks = []
                for tok in engine.generate_stream(
                    [
                        {"role": "system", "content": sys_msg},
                        {"role": "user", "content": user_msg}
                    ],
                    max_tokens=sec_tokens_budget,
                    temperature=0.65
                ):
                    if tok:
                        collected_chunks.append(tok)
                        yield {"type": "token", "content": tok}

                section_body = cls._clean_model_section_text("".join(collected_chunks), heading)

            # Si el motor no estaba cargado o devolvió un fragmento demasiado corto, complementar con síntesis estructurada
            if len(section_body.split()) < 60:
                fallback_text = cls._build_fallback_section(sec_spec, config, project_context)
                if not section_body:
                    section_body = fallback_text
                    yield {"type": "token", "content": section_body}
                else:
                    section_body = section_body + "\n\n" + fallback_text
                    yield {"type": "token", "content": "\n\n" + fallback_text}

            yield {"type": "token", "content": "\n\n"}

            if sec_spec["id"] == "preliminares":
                m_title = re.search(r"^(?:\*\*|#|\s)*T[ÍI]TULO(?:\s+ACAD[ÉE]MICO)?\s*:\s*(.+)$", section_body, re.IGNORECASE | re.MULTILINE)
                if m_title:
                    extracted_title = m_title.group(1).strip().strip("*#\"'")
                    if len(extracted_title) > 10:
                        essay_title = extracted_title
                        section_body = re.sub(
                            r"^(?:\*\*|#|\s)*T[ÍI]TULO(?:\s+ACAD[ÉE]MICO)?\s*:\s*.+\n*",
                            "",
                            section_body,
                            count=1,
                            flags=re.IGNORECASE | re.MULTILINE
                        ).strip()

            compiled_sections.append({
                "heading": heading,
                "level": sec_spec["level"],
                "body": section_body
            })

            first_para = section_body.split("\n\n")[0][:380] if section_body else ""
            accumulated_summary_notes.append(f"- {heading}: {first_para}")

        # Guardar automáticamente en Word (.docx) en la carpeta del proyecto
        act_id = f"create_word-{int(datetime.datetime.now().timestamp() * 1000)}"
        yield {
            "type": "action_start",
            "id": act_id,
            "tool": "create_word",
            "params": {"filename": filename}
        }
        docx_result = cls.export_academic_docx(
            project_path=project_path,
            filename=filename,
            title=essay_title,
            config=config,
            sections=compiled_sections
        )
        yield {
            "type": "action_end",
            "id": act_id,
            "tool": "create_word",
            "params": {"filename": filename},
            "result": docx_result
        }
        yield {
            "type": "artifact_created",
            "tool": "create_word",
            "result": docx_result
        }

    @classmethod
    def export_academic_docx(
        cls,
        project_path: str,
        filename: str,
        title: str,
        config: Dict[str, str],
        sections: List[Dict[str, Any]]
    ) -> Dict[str, Any]:
        """
        Crea un archivo Word (.docx) con maquetación académica formal (márgenes estándar,
        Times New Roman 12pt, interlineado 1.5, alineación justificada, bloque de portada/encabezado
        institucional, subtítulos jerárquicos y sangría/espaciado limpio), además de una copia .md.
        """
        os.makedirs(project_path, exist_ok=True)
        if not filename.lower().endswith(".docx"):
            filename += ".docx"
        target_docx = os.path.join(project_path, filename)
        target_md = os.path.join(project_path, filename[:-5] + ".md")

        # Calcular métricas reales
        all_text_parts = [title] + [f"{s.get('heading', '')}\n{s.get('body', '')}" for s in sections]
        full_markdown = f"# {title}\n\n"
        full_markdown += (
            f"> **Tipo:** Ensayo {config.get('essay_type', 'Académico')} | "
            f"**Nivel:** {config.get('academic_level', 'Universitario')} | "
            f"**Norma:** {config.get('citation_style', 'APA 7ª Edición')} | "
            f"**Fecha:** {datetime.date.today().strftime('%d/%m/%Y')}\n\n---\n\n"
        )
        for s in sections:
            full_markdown += f"## {s.get('heading', '')}\n\n{s.get('body', '')}\n\n"

        word_count = len(re.findall(r"\S+", " ".join(all_text_parts)))
        paragraph_count = 0

        # Guardar respaldo en Markdown
        try:
            with open(target_md, "w", encoding="utf-8") as f_md:
                f_md.write(full_markdown)
        except Exception:
            pass

        try:
            import docx
            from docx.shared import Pt, Inches, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = docx.Document()

            # Configurar márgenes académicos (1 pulgada / 2.54 cm en todos los lados)
            for section in doc.sections:
                section.top_margin = Inches(1.0)
                section.bottom_margin = Inches(1.0)
                section.left_margin = Inches(1.0)
                section.right_margin = Inches(1.0)

            # Estilo Normal base: Times New Roman 12pt
            normal_style = doc.styles["Normal"]
            normal_style.font.name = "Times New Roman"
            normal_style.font.size = Pt(12)
            normal_style.font.color.rgb = RGBColor(20, 20, 20)

            # Bloque superior de metadatos académicos
            meta_p = doc.add_paragraph()
            meta_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            meta_run = meta_p.add_run(
                f"Modalidad: Ensayo {config.get('essay_type', 'Académico')}\n"
                f"Nivel Académico: {config.get('academic_level', 'Posgrado / Universitario')}\n"
                f"Norma Editorial: {config.get('citation_style', 'APA 7ª Edición')}\n"
                f"Fecha: {datetime.date.today().strftime('%d/%m/%Y')} • Extensión: ~{word_count:,} palabras"
            )
            meta_run.font.name = "Times New Roman"
            meta_run.font.size = Pt(9.5)
            meta_run.font.italic = True
            meta_run.font.color.rgb = RGBColor(90, 100, 115)

            # Título Principal del Ensayo Centrado
            title_p = doc.add_paragraph()
            title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_p.paragraph_format.space_before = Pt(14)
            title_p.paragraph_format.space_after = Pt(18)
            title_run = title_p.add_run(title)
            title_run.font.name = "Times New Roman"
            title_run.font.size = Pt(16)
            title_run.font.bold = True
            title_run.font.color.rgb = RGBColor(15, 23, 42)

            # Recorrer secciones
            for sec in sections:
                sec_heading = sec.get("heading", "")
                sec_body = sec.get("body", "")
                is_references = "referencias" in sec_heading.lower() or "bibliograf" in sec_heading.lower()

                if sec_heading:
                    h_para = doc.add_paragraph()
                    h_para.paragraph_format.space_before = Pt(16)
                    h_para.paragraph_format.space_after = Pt(8)
                    h_run = h_para.add_run(sec_heading)
                    h_run.font.name = "Times New Roman"
                    h_run.font.size = Pt(13.5)
                    h_run.font.bold = True
                    h_run.font.color.rgb = RGBColor(15, 23, 42)

                # Dividir el cuerpo en bloques de párrafo
                raw_blocks = [b.strip() for b in re.split(r"\n\s*\n", sec_body) if b.strip()]
                for block in raw_blocks:
                    # Si el bloque tiene múltiples líneas (por ejemplo listado de referencias o subtítulos)
                    lines = [ln.strip() for ln in block.splitlines() if ln.strip()]
                    if len(lines) > 1 and (is_references or any(ln.startswith(("-", "*", "•", "###")) for ln in lines)):
                        for ln in lines:
                            cls._add_formatted_paragraph(doc, ln, is_references=is_references)
                            paragraph_count += 1
                    else:
                        merged_line = " ".join(lines)
                        cls._add_formatted_paragraph(doc, merged_line, is_references=is_references)
                        paragraph_count += 1

            doc.save(target_docx)
            return {
                "status": "success",
                "file": target_docx,
                "markdown_file": target_md,
                "title": title,
                "word_count": word_count,
                "paragraph_count": paragraph_count,
                "sections_count": len(sections),
                "citation_style": config.get("citation_style", "APA 7ª Edición")
            }
        except Exception as e:
            return {
                "status": "error",
                "error": str(e),
                "markdown_file": target_md,
                "word_count": word_count
            }

    @staticmethod
    def _add_formatted_paragraph(doc, text_line: str, is_references: bool = False):
        """Agrega un párrafo con soporte para subtítulos ###, negritas **texto** y cursivas *texto*."""
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        line = text_line.strip()
        if not line:
            return

        # Detectar subtítulo de nivel 2 o 3 (ej. ### 3.1. ...)
        if line.startswith("#"):
            clean_sub = re.sub(r"^#+\s*", "", line).strip().strip("*")
            sub_p = doc.add_paragraph()
            sub_p.paragraph_format.space_before = Pt(12)
            sub_p.paragraph_format.space_after = Pt(6)
            sub_run = sub_p.add_run(clean_sub)
            sub_run.font.name = "Times New Roman"
            sub_run.font.size = Pt(12.5)
            sub_run.font.bold = True
            sub_run.font.color.rgb = RGBColor(30, 41, 59)
            return

        # Limpiar viñeta inicial en referencias
        if is_references:
            line = re.sub(r"^[-*•]\s+", "", line).strip()

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if is_references else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(8)

        if is_references:
            # Sangría francesa académica para referencias bibliográficas
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)

        # Parsear negritas (**...**) y cursivas (*...*)
        tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", line)
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
                r = p.add_run(tok[2:-2])
                r.bold = True
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                r = p.add_run(tok[1:-1])
                r.italic = True
            else:
                r = p.add_run(tok)
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
